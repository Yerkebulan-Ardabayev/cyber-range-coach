from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import stat
import subprocess
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from fastapi import UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import Settings, project_root
from ..errors import AppError
from ..models import Draft, SourceBlock, StudioSource
from ..schemas import DraftResponse, StudioSourceResponse
from .commands import hidden_window_flags

ALLOWED_EXTENSIONS = {".docx", ".md", ".txt", ".log", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}


def tesseract_binary() -> Path | None:
    bundled_name = "tesseract.exe" if os.name == "nt" else "tesseract"
    bundled = project_root() / "vendor" / "tesseract" / bundled_name
    if bundled.is_file():
        return bundled
    discovered = shutil.which("tesseract")
    return Path(discovered) if discovered else None


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_filename(filename: str | None) -> str:
    value = Path(filename or "upload").name
    value = re.sub(r"[^A-Za-zА-Яа-яЁё0-9._ -]+", "_", value).strip(" .")
    return value[:200] or "upload"


async def import_source(
    session: Session,
    settings: Settings,
    upload: UploadFile,
) -> StudioSourceResponse:
    original_name = safe_filename(upload.filename)
    suffix = Path(original_name).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise AppError(
            415,
            "unsupported_source",
            "Studio поддерживает DOCX, Markdown, TXT, LOG, PNG, JPEG и TIFF.",
        )
    data = await upload.read(settings.max_import_bytes + 1)
    if len(data) > settings.max_import_bytes:
        raise AppError(413, "source_too_large", "Файл превышает локальный лимит Studio.")
    if not data:
        raise AppError(422, "empty_source", "Пустой файл нельзя импортировать.")
    digest = sha256_bytes(data)
    existing = session.scalar(select(StudioSource).where(StudioSource.sha256 == digest))
    if existing:
        return studio_source_response(session, existing)
    stored_name = f"{digest}{suffix}"
    destination = settings.imports_dir / stored_name
    destination.write_bytes(data)
    if sha256_file(destination) != digest:
        destination.unlink(missing_ok=True)
        raise AppError(
            500, "snapshot_hash_mismatch", "Не удалось подтвердить целостность snapshot."
        )
    blocks, extractor = extract_blocks(destination, suffix)
    destination.chmod(stat.S_IRUSR | stat.S_IRGRP | stat.S_IROTH)
    source = StudioSource(
        original_name=original_name,
        stored_name=stored_name,
        media_type=upload.content_type or "application/octet-stream",
        sha256=digest,
        size_bytes=len(data),
        extractor=extractor,
        immutable_ok=destination.stat().st_mode & 0o222 == 0,
    )
    session.add(source)
    session.flush()
    for ordinal, locator, text in blocks:
        session.add(SourceBlock(source_id=source.id, ordinal=ordinal, locator=locator, text=text))
    session.commit()
    session.refresh(source)
    if sha256_file(destination) != digest:
        raise AppError(500, "snapshot_changed", "Snapshot изменился после extraction.")
    return studio_source_response(session, source)


def extract_blocks(path: Path, suffix: str) -> tuple[list[tuple[int, str, str]], str]:
    if suffix in {".md", ".txt", ".log"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        blocks = [block.strip() for block in re.split(r"\n\s*\n", text) if block.strip()]
        return [(index, f"block:{index}", block) for index, block in enumerate(blocks, 1)], "text"
    if suffix == ".docx":
        document = Document(str(path))
        extracted: list[tuple[int, str, str]] = []
        ordinal = 1
        for paragraph_index, paragraph in enumerate(document.paragraphs, 1):
            text = paragraph.text.strip()
            if text:
                extracted.append((ordinal, f"paragraph:{paragraph_index}", text))
                ordinal += 1
        for table_index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                text = " | ".join(cell.text.strip() for cell in row.cells).strip(" |")
                if text:
                    extracted.append((ordinal, f"table:{table_index}:row:{row_index}", text))
                    ordinal += 1
        return extracted, "python-docx"
    if suffix in IMAGE_EXTENSIONS:
        executable = tesseract_binary()
        if not executable:
            return [], "ocr_unavailable"
        safe_env = {
            key: value
            for key, value in os.environ.items()
            if key in {"PATH", "PATHEXT", "SYSTEMROOT", "WINDIR", "TEMP", "TMP"}
        }
        bundled_tessdata = executable.parent / "tessdata"
        if bundled_tessdata.is_dir():
            safe_env["TESSDATA_PREFIX"] = str(bundled_tessdata)
        for languages in ("eng+rus", "eng"):
            result = subprocess.run(
                [str(executable), str(path), "stdout", "-l", languages],
                capture_output=True,
                text=True,
                timeout=60,
                check=False,
                env=safe_env,
                creationflags=hidden_window_flags(),
            )
            if result.returncode == 0 and result.stdout.strip():
                blocks = [
                    block.strip() for block in re.split(r"\n\s*\n", result.stdout) if block.strip()
                ]
                return (
                    [(index, f"ocr:block:{index}", block) for index, block in enumerate(blocks, 1)],
                    f"tesseract:{languages}",
                )
        return [], "ocr_failed"
    raise AppError(415, "unsupported_source", "Неизвестный тип source.")


def studio_source_response(session: Session, source: StudioSource) -> StudioSourceResponse:
    count = session.query(SourceBlock).filter(SourceBlock.source_id == source.id).count()
    return StudioSourceResponse.model_validate(
        {
            **{column.name: getattr(source, column.name) for column in source.__table__.columns},
            "block_count": count,
        }
    )


def create_draft(session: Session, title: str, source_ids: list[int]) -> Draft:
    sources = session.scalars(select(StudioSource).where(StudioSource.id.in_(source_ids))).all()
    if len(sources) != len(set(source_ids)):
        raise AppError(404, "source_not_found", "Один или несколько Studio sources не найдены.")
    blocks = session.scalars(
        select(SourceBlock)
        .where(SourceBlock.source_id.in_(source_ids))
        .order_by(SourceBlock.source_id, SourceBlock.ordinal)
    ).all()
    if not blocks:
        raise AppError(
            409,
            "source_has_no_text",
            "В источнике нет извлечённого текста. Проверьте OCR runtime или формат документа.",
        )
    sections: list[dict[str, Any]] = []
    for block in blocks:
        sections.append(
            {
                "heading": f"Источник {block.source_id}, {block.locator}",
                "body": block.text,
                "source_block_ids": [block.id],
            }
        )
    draft = Draft(
        title=title,
        source_ids=sorted(set(source_ids)),
        covered_block_ids=[block.id for block in blocks],
        content={
            "title": title,
            "objectives": ["Проверить и превратить source blocks в самостоятельный учебный цикл."],
            "sections": sections,
            "publication_note": "Черновик требует методического review владельцем.",
        },
    )
    session.add(draft)
    session.commit()
    session.refresh(draft)
    return draft


def validate_draft(session: Session, settings: Settings, draft: Draft) -> dict[str, Any]:
    sources = session.scalars(
        select(StudioSource).where(StudioSource.id.in_(draft.source_ids))
    ).all()
    blocks = session.scalars(
        select(SourceBlock).where(SourceBlock.source_id.in_(draft.source_ids))
    ).all()
    expected = {block.id for block in blocks}
    raw_sections = draft.content.get("sections")
    sections = raw_sections if isinstance(raw_sections, list) else []
    covered: set[int] = set()
    sections_ok = bool(sections)
    for section in sections:
        if not isinstance(section, dict):
            sections_ok = False
            continue
        raw_ids = section.get("source_block_ids")
        if not isinstance(raw_ids, list) or not raw_ids:
            sections_ok = False
            continue
        if not str(section.get("heading", "")).strip() or not str(section.get("body", "")).strip():
            sections_ok = False
        for block_id in raw_ids:
            if isinstance(block_id, int):
                covered.add(block_id)
            else:
                sections_ok = False
    missing = sorted(expected - covered)
    unknown = sorted(covered - expected)
    snapshot_checks: list[dict[str, Any]] = []
    snapshots_ok = True
    for source in sources:
        path = settings.imports_dir / source.stored_name
        actual = sha256_file(path) if path.exists() else None
        read_only = bool(path.exists() and path.stat().st_mode & 0o222 == 0)
        ok = actual == source.sha256 and read_only
        snapshots_ok = snapshots_ok and ok
        snapshot_checks.append(
            {
                "source_id": source.id,
                "expected_sha256": source.sha256,
                "actual_sha256": actual,
                "ok": ok,
                "read_only": read_only,
            }
        )
    title_ok = bool(str(draft.content.get("title", "")).strip())
    objectives_ok = bool(draft.content.get("objectives"))
    valid = (
        snapshots_ok and sections_ok and not missing and not unknown and title_ok and objectives_ok
    )
    report = {
        "valid": valid,
        "source_count": len(sources),
        "block_count": len(expected),
        "covered_count": len(expected & covered),
        "missing_block_ids": missing,
        "unknown_block_ids": unknown,
        "snapshots": snapshot_checks,
        "title_ok": title_ok,
        "objectives_ok": objectives_ok,
        "sections_ok": sections_ok,
    }
    draft.covered_block_ids = sorted(covered)
    draft.validation_report = report
    draft.status = "validated" if valid else "draft"
    session.commit()
    return report


def publish_draft(session: Session, settings: Settings, draft: Draft) -> Draft:
    if draft.status != "validated":
        raise AppError(
            409,
            "draft_not_validated",
            "Сначала сохраните review-правки и явно выполните Validate.",
        )
    report = validate_draft(session, settings, draft)
    if not report["valid"]:
        raise AppError(
            409, "draft_not_valid", "Draft не прошёл source coverage validation.", report
        )
    output = {
        "id": f"studio-draft-{draft.id}",
        "version": 1,
        "published_at": datetime.now(UTC).isoformat(),
        "source_ids": draft.source_ids,
        "source_block_ids": draft.covered_block_ids,
        "content": draft.content,
    }
    final_path = settings.published_content_dir / f"studio-draft-{draft.id}.json"
    temporary = final_path.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(final_path)
    final_path.chmod(stat.S_IRUSR | stat.S_IRGRP | stat.S_IROTH)
    draft.status = "published"
    draft.published_at = datetime.now(UTC)
    session.commit()
    session.refresh(draft)
    return draft


def update_draft(session: Session, draft: Draft, title: str, content: dict[str, Any]) -> Draft:
    if draft.status == "published":
        raise AppError(409, "published_draft_immutable", "Опубликованный draft неизменяем.")
    draft.title = title
    draft.content = content
    draft.status = "draft"
    draft.validation_report = {}
    session.commit()
    session.refresh(draft)
    return draft


def draft_response(draft: Draft) -> DraftResponse:
    return DraftResponse.model_validate(draft)
