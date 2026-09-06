#!/usr/bin/env python3
"""Extract addressable text blocks from a mnemonic .docx source.

The address scheme is the contract used by curriculum/commands/*.yaml:

* a paragraph is ``P%03d`` of its 1-based index among ALL paragraphs of the
  document, empty ones included, and only non-empty paragraphs are emitted;
* a table row is ``T%02dR%03d`` of its table and row index.

The empty paragraphs deliberately consume numbers: that is what makes an
address stable when a blank line is later filled in.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.table import Table  # noqa: F401


def extract(path: Path) -> list[dict[str, str]]:
    document = Document(str(path))
    blocks: list[dict[str, str]] = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            blocks.append({"kind": "paragraph", "address": f"P{index:03d}", "text": text})
    for table_index, table in enumerate(document.tables, 1):
        for row_index, row in enumerate(table.rows, 1):
            text = " | ".join(cell.text.strip() for cell in row.cells).strip(" |")
            if text:
                blocks.append(
                    {
                        "kind": "table_row",
                        "address": f"T{table_index:02d}R{row_index:03d}",
                        "text": text,
                    }
                )
    return blocks


def embedded_images(path: Path) -> int:
    import zipfile

    with zipfile.ZipFile(path) as archive:
        return sum(1 for name in archive.namelist() if name.startswith("word/media/"))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("paths", nargs="+", type=Path)
    parser.add_argument("--json", action="store_true", help="emit JSON instead of text")
    args = parser.parse_args()
    payload = []
    for path in args.paths:
        if not path.is_file():
            print(f"missing source: {path}", file=sys.stderr)
            return 1
        blocks = extract(path)
        payload.append(
            {
                "name": path.name,
                "block_count": len(blocks),
                "embedded_images": embedded_images(path),
                "blocks": blocks,
            }
        )
    if args.json:
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0
    for source in payload:
        print(f"=== {source['name']} blocks={source['block_count']} images={source['embedded_images']}")
        for block in source["blocks"]:
            print(f"{block['address']}\t{block['kind'][:1]}\t{block['text']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
